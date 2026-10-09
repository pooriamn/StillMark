
from __future__ import annotations

import copy
import json
from collections import defaultdict
from datetime import datetime, timezone
from pathlib import Path
from typing import Any

import yaml

from helpers_content import (
    CONTENT_DIR,
    ROOT,
    available_page_keys,
    available_series_slugs,
    available_work_ids,
    create_series_file,
    load_page_payload,
    load_yaml,
    page_file_for_key,
    reorder_homepage_featured,
    save_page_payload,
    series_file_for_slug,
    snapshot_path,
    work_file_for_id,
    work_to_series_map,
    write_yaml,
)
from helpers_image import (
    derivative_dir_for_work,
    generate_derivatives,
    load_pipeline,
    move_generated_between_series,
    move_original_between_series,
    source_path_for_work,
)
from og_images import ensure_og_images_from_content, load_content_for_og

WORKBOOK_VERSION = 'STILLMRK_WORKBOOK_V1'
EXPORT_DIR = ROOT / 'exports'


def _ensure_openpyxl():
    try:
        from openpyxl import Workbook, load_workbook  # noqa: F401
        from openpyxl.comments import Comment
        from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
        from openpyxl.worksheet.datavalidation import DataValidation
        from openpyxl.worksheet.table import Table, TableStyleInfo
        from openpyxl.utils import get_column_letter
    except ImportError as exc:
        raise RuntimeError(
            'This feature needs openpyxl. Install it with: python -m pip install openpyxl'
        ) from exc
    return {
        'Workbook': Workbook,
        'load_workbook': load_workbook,
        'Comment': Comment,
        'Alignment': Alignment,
        'Border': Border,
        'Font': Font,
        'PatternFill': PatternFill,
        'Side': Side,
        'DataValidation': DataValidation,
        'Table': Table,
        'TableStyleInfo': TableStyleInfo,
        'get_column_letter': get_column_letter,
    }


def _now_stamp() -> str:
    return datetime.now(timezone.utc).strftime('%Y%m%dT%H%M%SZ')


def _normalize_sheet_name(name: str) -> str:
    safe = name.replace('-', '_').replace(' ', '_')
    safe = ''.join(ch for ch in safe if ch.isalnum() or ch in {'_', ' '})
    return safe[:31]


def _path_tokens(path: str) -> list[Any]:
    tokens: list[Any] = []
    token = ''
    i = 0
    while i < len(path):
        ch = path[i]
        if ch == '.':
            if token:
                tokens.append(token)
                token = ''
            i += 1
            continue
        if ch == '[':
            if token:
                tokens.append(token)
                token = ''
            j = path.index(']', i)
            tokens.append(int(path[i + 1:j]) - 1)
            i = j + 1
            continue
        token += ch
        i += 1
    if token:
        tokens.append(token)
    return tokens


def _flatten_value(value: Any, prefix: str = '') -> list[dict[str, Any]]:
    rows: list[dict[str, Any]] = []
    if isinstance(value, dict):
        if not value and prefix:
            rows.append({'path': prefix, 'value': '', 'type': 'empty_object'})
        for key, item in value.items():
            next_prefix = f'{prefix}.{key}' if prefix else str(key)
            rows.extend(_flatten_value(item, next_prefix))
        return rows
    if isinstance(value, list):
        if not value:
            rows.append({'path': prefix, 'value': '', 'type': 'empty_list'})
            return rows
        for index, item in enumerate(value, start=1):
            next_prefix = f'{prefix}[{index}]'
            rows.extend(_flatten_value(item, next_prefix))
        return rows
    value_type = 'string'
    if isinstance(value, bool):
        value_type = 'bool'
    elif isinstance(value, int):
        value_type = 'int'
    elif isinstance(value, float):
        value_type = 'float'
    elif value is None:
        value_type = 'none'
    rows.append({'path': prefix, 'value': value, 'type': value_type})
    return rows


def _set_nested_value(target: Any, tokens: list[Any], value: Any) -> Any:
    if not tokens:
        return value
    head, *tail = tokens
    if isinstance(head, str):
        if not isinstance(target, dict):
            target = {}
        target[head] = _set_nested_value(target.get(head), tail, value)
        return target
    # list index
    if not isinstance(target, list):
        target = []
    while len(target) <= head:
        target.append(None)
    target[head] = _set_nested_value(target[head], tail, value)
    return target


def _unflatten_rows(rows: list[dict[str, Any]]) -> Any:
    result: Any = {}
    for row in rows:
        path = str(row.get('path') or '').strip()
        if not path:
            continue
        row_type = str(row.get('type') or 'string')
        raw_value = row.get('value')
        if row_type == 'empty_list':
            value: Any = []
        elif row_type == 'empty_object':
            value = {}
        elif row_type == 'bool':
            if isinstance(raw_value, bool):
                value = raw_value
            else:
                value = str(raw_value).strip().lower() in {'true', '1', 'yes', 'y'}
        elif row_type == 'int':
            value = int(raw_value) if raw_value not in {None, ''} else 0
        elif row_type == 'float':
            value = float(raw_value) if raw_value not in {None, ''} else 0.0
        elif row_type == 'none':
            value = None
        else:
            value = '' if raw_value is None else raw_value
        result = _set_nested_value(result, _path_tokens(path), value)
    return result


def _site_settings_rows() -> list[dict[str, Any]]:
    site = load_yaml(CONTENT_DIR / 'site.yaml') or {}
    artist = load_yaml(CONTENT_DIR / 'artist.yaml') or {}
    rows: list[dict[str, Any]] = []
    for group, payload in [('site', site), ('artist', artist)]:
        for key, value in payload.items():
            rows.append({
                'group': group,
                'key': key,
                'value': json.dumps(value, ensure_ascii=False) if isinstance(value, (list, dict)) else value,
                'type': 'json' if isinstance(value, (list, dict)) else ('bool' if isinstance(value, bool) else 'string'),
            })
    return rows


def _document_rows() -> list[dict[str, Any]]:
    resources = load_yaml(CONTENT_DIR / 'resources.yaml') or {}
    rows = []
    for item in resources.get('downloads') or []:
        rows.append({
            'id': item.get('id', ''),
            'title': item.get('title', ''),
            'kind': item.get('kind', ''),
            'description': item.get('description', ''),
            'file': item.get('file', ''),
            'audience': item.get('audience', ''),
            'featured': bool(item.get('featured', False)),
        })
    return rows


def _navigation_rows() -> list[dict[str, Any]]:
    nav = load_yaml(CONTENT_DIR / 'navigation.yaml') or {}
    rows = []
    for idx, item in enumerate(nav.get('items') or [], start=1):
        rows.append({
            'position': idx,
            'label': item.get('label', ''),
            'href': item.get('href', ''),
            'page': item.get('page', ''),
            'visible': bool(item.get('visible', True)) if 'visible' in item else True,
        })
    return rows


def _page_rows(page_key: str) -> list[dict[str, Any]]:
    payload = copy.deepcopy(load_page_payload(page_key))
    if page_key == 'home':
        fs = payload.get('featured_series') if isinstance(payload.get('featured_series'), dict) else {}
        fs.pop('series_slugs', None)
        sw = payload.get('selected_works') if isinstance(payload.get('selected_works'), dict) else {}
        sw.pop('work_ids', None)
    rows = _flatten_value(payload)
    for row in rows:
        row['editable'] = True
        row['notes'] = ''
    return rows


def _series_rows() -> list[dict[str, Any]]:
    rows: list[dict[str, Any]] = []
    for slug in available_series_slugs():
        payload = load_yaml(series_file_for_slug(slug)) or {}
        rows.append({
            'series_slug': slug,
            'title': payload.get('title', ''),
            'years': payload.get('years', ''),
            'mood': payload.get('mood', ''),
            'description': payload.get('description', ''),
            'cover_work_id': payload.get('cover_work_id', ''),
            'order': payload.get('order', ''),
            'visibility': payload.get('visibility', 'public'),
            'review_mode': bool(payload.get('review_mode', False)),
            'allow_favorites': bool(payload.get('allow_favorites', True)),
            'allow_inquiry_basket': bool(payload.get('allow_inquiry_basket', True)),
            'project_type': payload.get('project_type', ''),
            'download_ids': ' | '.join(str(x) for x in (payload.get('download_ids') or [])),
        })
    return rows


def _series_position_map() -> dict[str, tuple[str, int]]:
    mapping: dict[str, tuple[str, int]] = {}
    for slug in available_series_slugs():
        payload = load_yaml(series_file_for_slug(slug)) or {}
        for idx, work_id in enumerate(payload.get('work_ids') or [], start=1):
            mapping[str(work_id)] = (slug, idx)
    return mapping


def _derivative_status_for_work(series_slug: str, work_id: str, render_name: str, pipeline: dict[str, Any] | None = None) -> str:
    """Return a cheap derivative readiness hint for workbook round-trips."""
    series_slug = str(series_slug or "").strip()
    work_id = str(work_id or "").strip()
    render_name = str(render_name or work_id).strip()
    if not series_slug or not work_id or not render_name:
        return "missing-reference"
    try:
        pipeline = pipeline or load_pipeline()
        source = source_path_for_work(series_slug, work_id, pipeline=pipeline)
        directory = derivative_dir_for_work(series_slug, render_name, pipeline=pipeline)
        derivative_files = []
        if directory.exists():
            derivative_files = [
                path for path in directory.iterdir()
                if path.is_file() and path.suffix.lower() in {".jpg", ".jpeg", ".webp"}
            ]
        if not source:
            return "source-missing"
        if not derivative_files:
            return "derivatives-missing"
        return "ready"
    except Exception:
        return "check"


def _work_rows() -> list[dict[str, Any]]:
    positions = _series_position_map()
    rows: list[dict[str, Any]] = []
    for work_id in available_work_ids():
        payload = load_yaml(work_file_for_id(work_id)) or {}
        series_slug, series_position = positions.get(work_id, (str(payload.get('series') or ''), 9999))
        image = payload.get('image') if isinstance(payload.get('image'), dict) else {}
        render_name = image.get('render_name', work_id) if image else work_id
        rows.append({
            'work_id': work_id,
            'title': payload.get('title', ''),
            'series_slug': series_slug,
            'series_position': series_position,
            'published': bool(payload.get('published', True)),
            'year': payload.get('year', ''),
            'location': payload.get('location', ''),
            'alt': payload.get('alt', ''),
            'caption': payload.get('caption', ''),
            'tags': ' | '.join(str(x) for x in (payload.get('tags') or [])),
            'hero_safe': bool(payload.get('hero_safe', False)),
            'grid_safe': bool(payload.get('grid_safe', False)),
            'social_safe': bool(payload.get('social_safe', False)),
            'focal_x': (payload.get('focal_point') or {}).get('x', 50),
            'focal_y': (payload.get('focal_point') or {}).get('y', 50),
            'project_type': payload.get('project_type', ''),
            'licensing_available': bool(payload.get('licensing_available', False)),
            'print_available': bool(payload.get('print_available', False)),
            'price_note': payload.get('price_note', ''),
            'image_master': image.get('master', payload.get('image', '')) if image else payload.get('image', ''),
            'render_name': render_name,
            'derivative_status': _derivative_status_for_work(series_slug, work_id, render_name),
            'status': 'OK' if payload.get('alt') and payload.get('title') else 'CHECK',
        })
    rows.sort(key=lambda row: (str(row['series_slug']), int(row['series_position']), str(row['work_id'])))
    return rows


def _home_order_rows() -> list[dict[str, Any]]:
    payload = load_page_payload('home')
    rows: list[dict[str, Any]] = []
    featured = payload.get('featured_series') if isinstance(payload.get('featured_series'), dict) else {}
    for idx, item in enumerate(featured.get('series_slugs') or [], start=1):
        rows.append({'group': 'featured_series', 'item_id': str(item), 'position': idx, 'enabled': True})
    selected = payload.get('selected_works') if isinstance(payload.get('selected_works'), dict) else {}
    for idx, item in enumerate(selected.get('work_ids') or [], start=1):
        rows.append({'group': 'selected_works', 'item_id': str(item), 'position': idx, 'enabled': True})
    return rows


def _page_preview_rows() -> list[dict[str, Any]]:
    rows = []
    for page_key in available_page_keys():
        payload = load_page_payload(page_key)
        meta = payload.get('meta') if isinstance(payload.get('meta'), dict) else {}
        hero = payload.get('hero') if isinstance(payload.get('hero'), dict) else {}
        featured = payload.get('featured_series') if isinstance(payload.get('featured_series'), dict) else {}
        selected = payload.get('selected_works') if isinstance(payload.get('selected_works'), dict) else {}
        rows.append({
            'page': page_key,
            'page_title': hero.get('title', payload.get('hero_title', '')),
            'hero_work_id': hero.get('feature_work_id', ''),
            'featured_series': ', '.join(str(x) for x in (featured.get('series_slugs') or [])),
            'selected_works': ', '.join(str(x) for x in (selected.get('work_ids') or [])),
            'og_image': meta.get('og_image', ''),
            'status': 'OK',
        })
    return rows


def _hero_cover_rows() -> list[dict[str, Any]]:
    work_map = {row['work_id']: row for row in _work_rows()}
    rows = []
    for page_key in available_page_keys():
        payload = load_page_payload(page_key)
        hero = payload.get('hero') if isinstance(payload.get('hero'), dict) else {}
        work_id = str(hero.get('feature_work_id') or '').strip()
        work = work_map.get(work_id, {})
        rows.append({
            'target_type': 'page',
            'target_id': page_key,
            'image_work_id': work_id,
            'title': work.get('title', ''),
            'hero_safe': work.get('hero_safe', ''),
            'social_safe': work.get('social_safe', ''),
            'alt_ok': 'OK' if len(str(work.get('alt', '')).split()) >= 5 else 'CHECK',
            'published': work.get('published', ''),
            'notes': '',
        })
    for row in _series_rows():
        work = work_map.get(str(row['cover_work_id']), {})
        rows.append({
            'target_type': 'series',
            'target_id': row['series_slug'],
            'image_work_id': row['cover_work_id'],
            'title': work.get('title', ''),
            'hero_safe': work.get('hero_safe', ''),
            'social_safe': work.get('social_safe', ''),
            'alt_ok': 'OK' if len(str(work.get('alt', '')).split()) >= 5 else 'CHECK',
            'published': work.get('published', ''),
            'notes': '',
        })
    return rows


def _series_helper_rows(series_slug: str) -> tuple[dict[str, Any], list[dict[str, Any]]]:
    series_payload = load_yaml(series_file_for_slug(series_slug)) or {}
    works = []
    for row in _work_rows():
        if row['series_slug'] == series_slug:
            works.append(row)
    overview = {
        'series_slug': series_slug,
        'title': series_payload.get('title', ''),
        'description': series_payload.get('description', ''),
        'cover_work_id': series_payload.get('cover_work_id', ''),
        'visibility': series_payload.get('visibility', 'public'),
        'count': len(works),
    }
    return overview, works


def _validation_rows() -> list[dict[str, Any]]:
    rows: list[dict[str, Any]] = []
    work_ids = set(available_work_ids())
    series_slugs = set(available_series_slugs())
    for row in _work_rows():
        if row['series_slug'] not in series_slugs:
            rows.append({'level': 'ERROR', 'area': 'work', 'id': row['work_id'], 'issue': f"Unknown series_slug: {row['series_slug']}"})
        if len(str(row['alt']).split()) < 5:
            rows.append({'level': 'WARN', 'area': 'work', 'id': row['work_id'], 'issue': 'Alt text has fewer than 5 words'})
        if not row['title']:
            rows.append({'level': 'WARN', 'area': 'work', 'id': row['work_id'], 'issue': 'Missing title'})
    for row in _series_rows():
        if row['cover_work_id'] and row['cover_work_id'] not in work_ids:
            rows.append({'level': 'ERROR', 'area': 'series', 'id': row['series_slug'], 'issue': f"Unknown cover_work_id: {row['cover_work_id']}"})
    preview_rows = _page_preview_rows()
    for row in preview_rows:
        hero_id = row['hero_work_id']
        if hero_id and hero_id not in work_ids:
            rows.append({'level': 'ERROR', 'area': 'page', 'id': row['page'], 'issue': f"Unknown hero_work_id: {hero_id}"})
    if not rows:
        rows.append({'level': 'OK', 'area': 'workbook', 'id': 'all', 'issue': 'No validation problems found'})
    return rows


def _missing_metadata_rows() -> list[dict[str, Any]]:
    rows = []
    for row in _work_rows():
        issues = []
        if not row['title']:
            issues.append('title')
        if len(str(row['alt']).split()) < 5:
            issues.append('alt')
        if not row['caption']:
            issues.append('caption')
        if issues:
            rows.append({'work_id': row['work_id'], 'series_slug': row['series_slug'], 'missing_or_weak': ', '.join(issues)})
    return rows


def _unpublished_rows() -> list[dict[str, Any]]:
    return [row for row in _work_rows() if not bool(row['published'])]



WORKBOOK_COLORS = {
    'ink': '0F172A',
    'ink_soft': '334155',
    'line': 'CBD5E1',
    'paper': 'FFFFFF',
    'paper_soft': 'F8FAFC',
    'paper_alt': 'F1F5F9',
    'source_header': '111827',
    'helper_header': '1D4ED8',
    'report_header': '7C2D12',
    'settings_header': '0F766E',
    'page_header': '4338CA',
    'warning': 'FEF3C7',
    'warning_text': '92400E',
    'error': 'FEE2E2',
    'error_text': '991B1B',
    'success': 'DCFCE7',
    'success_text': '166534',
    'accent': 'E2E8F0',
}

HEADER_NOTES = {
    'work_id': 'Stable internal identifier. Do not rename casually.',
    'series_slug': 'Must match an existing series slug exactly.',
    'series_position': 'Lower number appears earlier inside the series.',
    'published': 'TRUE shows the work publicly. FALSE keeps it in the archive only.',
    'alt': 'Describe the visible content clearly. Use at least 5 words.',
    'caption': 'Short public-facing text. Avoid placeholder wording.',
    'tags': 'Separate multiple tags with a vertical bar: tag1 | tag2 | tag3',
    'hero_work_id': 'Must match an existing work_id. Used for the page hero image.',
    'cover_work_id': 'Must match an existing work_id inside that series.',
    'visible': 'TRUE shows this item in navigation. FALSE hides it.',
    'enabled': 'TRUE keeps the item active in homepage ordering. FALSE removes it from output.',
    'focal_x': 'Horizontal focal point as a percentage from 0 to 100.',
    'focal_y': 'Vertical focal point as a percentage from 0 to 100.',
    'og_image': 'Generated social-preview image path used by the public page.',
    'value': 'Edit only values you understand. Import will validate and preview changes.',
    'type': 'Controls how the importer interprets the value column.',
    'status': 'Calculated export status only. Review before importing.',
}

SHEET_TABS = {
    'README': '0F172A',
    'SITE_SETTINGS': '0F766E',
    'NAVIGATION': '0F766E',
    'DOCUMENTS': '0F766E',
    'PAGES_SUMMARY': '1D4ED8',
    'PAGE_PREVIEW_MAP': '1D4ED8',
    'HERO_AND_COVERS': '1D4ED8',
    'VALIDATION_REPORT': '7C2D12',
    'SERIES_MASTER': '4338CA',
    'WORKS_MASTER': '4338CA',
    'HOME_ORDER': '4338CA',
    'UNPUBLISHED': '7C3AED',
    'MISSING_METADATA': 'B45309',
}


def _sheet_kind_for_name(sheet_name: str) -> str:
    if sheet_name == 'README':
        return 'readme'
    if sheet_name in {'SITE_SETTINGS', 'NAVIGATION', 'DOCUMENTS'}:
        return 'settings'
    if sheet_name.startswith('PAGE__') or sheet_name in {'SERIES_MASTER', 'WORKS_MASTER', 'HOME_ORDER'}:
        return 'source'
    if sheet_name in {'VALIDATION_REPORT', 'MISSING_METADATA'}:
        return 'report'
    return 'helper'


def _set_tab_color(ws) -> None:
    color = SHEET_TABS.get(ws.title)
    if not color and ws.title.startswith('PAGE__'):
        color = WORKBOOK_COLORS['page_header']
    if not color and ws.title.startswith('SERIES__'):
        color = '6366F1'
    if not color:
        return
    try:
        ws.sheet_properties.tabColor = color
    except Exception:
        pass


def _safe_table_name(sheet_name: str, start_row: int) -> str:
    base = ''.join(ch if ch.isalnum() else '_' for ch in sheet_name)
    if not base or not base[0].isalpha():
        base = f'T_{base}'
    return f'{base[:22]}_{start_row}'


def _column_width_hint(header: str) -> int | None:
    hints = {
        'path': 34,
        'value': 42,
        'notes': 28,
        'description': 42,
        'meta_description': 38,
        'og_description': 38,
        'alt': 38,
        'caption': 38,
        'tags': 28,
        'file': 34,
        'href': 28,
        'og_image': 32,
        'featured_series': 30,
        'selected_works': 34,
        'issue': 42,
        'missing_or_weak': 28,
    }
    return hints.get(header)


def _apply_header_comments(ws, headers: list[str], header_row: int, styles: dict[str, Any]) -> None:
    Comment = styles.get('Comment')
    if Comment is None:
        return
    for idx, header in enumerate(headers, start=1):
        note = HEADER_NOTES.get(str(header).strip())
        if not note:
            continue
        ws.cell(header_row, idx).comment = Comment(note, 'STILLMRK Studio')


def _write_table(ws, headers: list[str], rows: list[dict[str, Any]], *, styles: dict[str, Any], freeze: str = 'A2', start_row: int = 1, sheet_kind: str | None = None, add_table: bool = True):
    Alignment = styles['Alignment']; Border = styles['Border']; Font = styles['Font']; PatternFill = styles['PatternFill']; Side = styles['Side']
    Table = styles.get('Table'); TableStyleInfo = styles.get('TableStyleInfo')
    sheet_kind = sheet_kind or _sheet_kind_for_name(ws.title)
    header_colors = {
        'source': WORKBOOK_COLORS['source_header'],
        'settings': WORKBOOK_COLORS['settings_header'],
        'helper': WORKBOOK_COLORS['helper_header'],
        'report': WORKBOOK_COLORS['report_header'],
        'readme': WORKBOOK_COLORS['ink'],
    }
    header_fill = PatternFill('solid', fgColor=header_colors.get(sheet_kind, WORKBOOK_COLORS['source_header']))
    header_font = Font(color='FFFFFF', bold=True, size=11)
    thin = Side(style='thin', color=WORKBOOK_COLORS['line'])
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    band_1 = PatternFill('solid', fgColor=WORKBOOK_COLORS['paper'])
    band_2 = PatternFill('solid', fgColor=WORKBOOK_COLORS['paper_soft'])
    warn_fill = PatternFill('solid', fgColor=WORKBOOK_COLORS['warning'])
    err_fill = PatternFill('solid', fgColor=WORKBOOK_COLORS['error'])
    ok_fill = PatternFill('solid', fgColor=WORKBOOK_COLORS['success'])
    header_row = start_row
    for col_idx, header in enumerate(headers, start=1):
        cell = ws.cell(header_row, col_idx, header)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
        cell.border = border
    for row_idx, row in enumerate(rows, start=header_row + 1):
        values = [row.get(h, '') for h in headers]
        for col_idx, value in enumerate(values, start=1):
            cell = ws.cell(row_idx, col_idx, value)
            cell.border = border
            cell.alignment = Alignment(vertical='top', wrap_text=True)
            cell.fill = band_1 if (row_idx - header_row) % 2 else band_2
            header_name = str(headers[col_idx - 1])
            if header_name in {'position', 'series_position', 'order', 'focal_x', 'focal_y'} and value not in {None, ''}:
                cell.alignment = Alignment(horizontal='center', vertical='top')
            if header_name in {'published', 'hero_safe', 'grid_safe', 'social_safe', 'review_mode', 'allow_favorites', 'allow_inquiry_basket', 'featured', 'enabled', 'visible', 'licensing_available', 'print_available'}:
                cell.alignment = Alignment(horizontal='center', vertical='top')
        row_text = ' | '.join(str(v) for v in values if v not in {None, ''}).upper()
        if 'ERROR' in row_text or 'UNKNOWN' in row_text:
            for cell in ws[row_idx]:
                cell.fill = err_fill
                cell.font = Font(color=WORKBOOK_COLORS['error_text'])
        elif 'WARN' in row_text or 'CHECK' in row_text:
            for cell in ws[row_idx]:
                cell.fill = warn_fill
                cell.font = Font(color=WORKBOOK_COLORS['warning_text'])
        elif 'OK' in row_text:
            for cell in ws[row_idx]:
                cell.fill = ok_fill
                cell.font = Font(color=WORKBOOK_COLORS['success_text'])
    end_row = max(header_row, header_row + len(rows))
    end_col = len(headers)
    end_letter = styles['get_column_letter'](end_col)
    ws.freeze_panes = freeze
    ws.auto_filter.ref = f"A{header_row}:{end_letter}{end_row}"
    ws.sheet_view.zoomScale = 90
    ws.sheet_view.showGridLines = sheet_kind != 'readme'
    ws.row_dimensions[header_row].height = 24
    for idx, header in enumerate(headers, start=1):
        header_name = str(header)
        max_len = max(12, len(header_name) + 2)
        width_hint = _column_width_hint(header_name)
        if width_hint:
            max_len = width_hint
        else:
            for row_idx in range(header_row + 1, end_row + 1):
                value = ws.cell(row_idx, idx).value
                if value is None:
                    continue
                max_len = min(max(max_len, len(str(value)) + 2), 48)
        ws.column_dimensions[styles['get_column_letter'](idx)].width = max_len
    if False and add_table and Table and TableStyleInfo and end_row > header_row:
        try:
            table = Table(displayName=_safe_table_name(ws.title, header_row), ref=f"A{header_row}:{end_letter}{end_row}")
            table_style = TableStyleInfo(
                name='TableStyleMedium2' if sheet_kind in {'source', 'settings'} else 'TableStyleMedium9',
                showFirstColumn=False,
                showLastColumn=False,
                showRowStripes=True,
                showColumnStripes=False,
            )
            table.tableStyleInfo = table_style
            ws.add_table(table)
        except Exception:
            pass
    _apply_header_comments(ws, headers, header_row, styles)
    _set_tab_color(ws)


def _add_bool_validation(ws, styles: dict[str, Any], column_letter: str, start_row: int = 2, end_row: int = 5000):
    return


def _build_readme_sheet(ws, styles: dict[str, Any]) -> None:
    Alignment = styles['Alignment']; Font = styles['Font']; PatternFill = styles['PatternFill']; Border = styles['Border']; Side = styles['Side']
    ws.sheet_view.showGridLines = False
    ws.merge_cells('A1:F2')
    ws['A1'] = 'STILLMRK Site Workbook'
    ws['A1'].font = Font(size=18, bold=True, color='FFFFFF')
    ws['A1'].fill = PatternFill('solid', fgColor=WORKBOOK_COLORS['ink'])
    ws['A1'].alignment = Alignment(horizontal='left', vertical='center')
    ws.row_dimensions[1].height = 28
    ws.row_dimensions[2].height = 28

    ws.merge_cells('A3:F4')
    ws['A3'] = 'Edit the source sheets, review the preview sheets, then import through STILLMRK Studio so validation, backups, and OG refresh run automatically.'
    ws['A3'].font = Font(size=11, color=WORKBOOK_COLORS['ink'])
    ws['A3'].fill = PatternFill('solid', fgColor=WORKBOOK_COLORS['paper_alt'])
    ws['A3'].alignment = Alignment(wrap_text=True, vertical='center')
    ws.row_dimensions[3].height = 24
    ws.row_dimensions[4].height = 24

    legend_rows = [
        {'Status colour': 'Green', 'Meaning': 'Looks healthy'},
        {'Status colour': 'Yellow', 'Meaning': 'Review before import'},
        {'Status colour': 'Red', 'Meaning': 'Broken or risky reference'},
    ]
    ws['A6'] = 'Quick start'
    ws['A6'].font = Font(bold=True, size=12, color=WORKBOOK_COLORS['ink'])
    quick_rows = [
        {'Step': '1', 'Action': 'Edit a source sheet', 'Where': 'WORKS_MASTER, SERIES_MASTER, PAGE__*, HOME_ORDER'},
        {'Step': '2', 'Action': 'Review likely outcomes', 'Where': 'PAGE_PREVIEW_MAP, HERO_AND_COVERS, VALIDATION_REPORT'},
        {'Step': '3', 'Action': 'Import with the Studio menu', 'Where': 'STILLMRK Studio > Import site workbook (.xlsx)'},
    ]
    _write_table(ws, ['Step', 'Action', 'Where'], quick_rows, styles=styles, freeze='A8', start_row=7, sheet_kind='readme', add_table=True)

    ws['A13'] = 'Workbook map'
    ws['A13'].font = Font(bold=True, size=12, color=WORKBOOK_COLORS['ink'])
    sheet_rows = [
        {'Sheet': 'SITE_SETTINGS', 'Purpose': 'Global site and artist settings', 'Edit': 'Yes'},
        {'Sheet': 'NAVIGATION', 'Purpose': 'Main site navigation labels and visibility', 'Edit': 'Yes'},
        {'Sheet': 'DOCUMENTS', 'Purpose': 'Public document cards and file links', 'Edit': 'Yes'},
        {'Sheet': 'PAGE__home / about / contact / portfolio / series', 'Purpose': 'Page content values', 'Edit': 'Yes'},
        {'Sheet': 'SERIES_MASTER', 'Purpose': 'Series-level metadata and cover work', 'Edit': 'Yes'},
        {'Sheet': 'WORKS_MASTER', 'Purpose': 'Main image/work metadata sheet', 'Edit': 'Yes'},
        {'Sheet': 'HOME_ORDER', 'Purpose': 'Homepage featured series and selected works order', 'Edit': 'Yes'},
        {'Sheet': 'PAGE_PREVIEW_MAP', 'Purpose': 'Likely public outcome by page', 'Edit': 'Review only'},
        {'Sheet': 'HERO_AND_COVERS', 'Purpose': 'Hero and cover image health check', 'Edit': 'Review only'},
        {'Sheet': 'VALIDATION_REPORT', 'Purpose': 'Warnings and broken references', 'Edit': 'Review only'},
    ]
    _write_table(ws, ['Sheet', 'Purpose', 'Edit'], sheet_rows, styles=styles, freeze='A14', start_row=14, sheet_kind='helper', add_table=True)

    ws['E6'] = 'Legend'
    ws['E6'].font = Font(bold=True, size=12, color=WORKBOOK_COLORS['ink'])
    _write_table(ws, ['Status colour', 'Meaning'], legend_rows, styles=styles, freeze='E8', start_row=7, sheet_kind='helper', add_table=True)
    ws.column_dimensions['A'].width = 18
    ws.column_dimensions['B'].width = 34
    ws.column_dimensions['C'].width = 34
    ws.column_dimensions['E'].width = 18
    ws.column_dimensions['F'].width = 24
    _set_tab_color(ws)


def export_site_workbook(output_path: str | Path | None = None) -> Path:
    ops = _ensure_openpyxl()
    Workbook = ops['Workbook']
    wb = Workbook()
    default = wb.active
    default.title = 'README'
    _build_readme_sheet(default, ops)

    site_ws = wb.create_sheet('SITE_SETTINGS')
    _write_table(site_ws, ['group', 'key', 'value', 'type'], _site_settings_rows(), styles=ops, sheet_kind='settings')

    nav_ws = wb.create_sheet('NAVIGATION')
    _write_table(nav_ws, ['position', 'label', 'href', 'page', 'visible'], _navigation_rows(), styles=ops, sheet_kind='settings')
    _add_bool_validation(nav_ws, ops, 'E')

    docs_ws = wb.create_sheet('DOCUMENTS')
    _write_table(docs_ws, ['id', 'title', 'kind', 'description', 'file', 'audience', 'featured'], _document_rows(), styles=ops, sheet_kind='settings')
    _add_bool_validation(docs_ws, ops, 'G')

    pages_summary = wb.create_sheet('PAGES_SUMMARY')
    _write_table(pages_summary, ['page', 'hero_title', 'hero_work_id', 'meta_title', 'meta_description'], [
        {
            'page': row['page'],
            'hero_title': row['page_title'],
            'hero_work_id': row['hero_work_id'],
            'meta_title': (load_page_payload(row['page']).get('meta') or {}).get('title', ''),
            'meta_description': (load_page_payload(row['page']).get('meta') or {}).get('description', ''),
        } for row in _page_preview_rows()
    ], styles=ops, sheet_kind='helper')

    for page_key in available_page_keys():
        ws = wb.create_sheet(f'PAGE__{page_key}')
        _write_table(ws, ['path', 'value', 'type', 'editable', 'notes'], _page_rows(page_key), styles=ops, sheet_kind='source')
        _add_bool_validation(ws, ops, 'D')

    series_master = wb.create_sheet('SERIES_MASTER')
    _write_table(series_master, ['series_slug', 'title', 'years', 'mood', 'description', 'cover_work_id', 'order', 'visibility', 'review_mode', 'allow_favorites', 'allow_inquiry_basket', 'project_type', 'download_ids'], _series_rows(), styles=ops, sheet_kind='source')
    _add_bool_validation(series_master, ops, 'I')
    _add_bool_validation(series_master, ops, 'J')
    _add_bool_validation(series_master, ops, 'K')

    works_master = wb.create_sheet('WORKS_MASTER')
    work_headers = ['work_id', 'title', 'series_slug', 'series_position', 'published', 'year', 'location', 'alt', 'caption', 'tags', 'hero_safe', 'grid_safe', 'social_safe', 'focal_x', 'focal_y', 'project_type', 'licensing_available', 'print_available', 'price_note', 'image_master', 'render_name', 'derivative_status', 'status']
    _write_table(works_master, work_headers, _work_rows(), styles=ops, sheet_kind='source')
    for col in ['E', 'K', 'L', 'M', 'Q', 'R']:
        _add_bool_validation(works_master, ops, col)

    home_ws = wb.create_sheet('HOME_ORDER')
    _write_table(home_ws, ['group', 'item_id', 'position', 'enabled'], _home_order_rows(), styles=ops, sheet_kind='source')
    _add_bool_validation(home_ws, ops, 'D')

    preview_ws = wb.create_sheet('PAGE_PREVIEW_MAP')
    _write_table(preview_ws, ['page', 'page_title', 'hero_work_id', 'featured_series', 'selected_works', 'og_image', 'status'], _page_preview_rows(), styles=ops, sheet_kind='helper')

    hc_ws = wb.create_sheet('HERO_AND_COVERS')
    _write_table(hc_ws, ['target_type', 'target_id', 'image_work_id', 'title', 'hero_safe', 'social_safe', 'alt_ok', 'published', 'notes'], _hero_cover_rows(), styles=ops, sheet_kind='helper')

    for slug in available_series_slugs():
        ws = wb.create_sheet(_normalize_sheet_name(f'SERIES__{slug}'))
        overview, works = _series_helper_rows(slug)
        ws.merge_cells('A1:F1')
        ws['A1'] = f"Series helper · {overview['title'] or slug}"
        ws['A1'].font = ops['Font'](size=14, bold=True, color='FFFFFF')
        ws['A1'].fill = ops['PatternFill']('solid', fgColor='6366F1')
        ws['A1'].alignment = ops['Alignment'](horizontal='left', vertical='center')
        ws['A2'] = 'Description'
        ws['B2'] = overview['description']
        ws['A3'] = 'Cover work'
        ws['B3'] = overview['cover_work_id']
        ws['A4'] = 'Visibility'
        ws['B4'] = overview['visibility']
        for cell in ['A2', 'A3', 'A4']:
            ws[cell].font = ops['Font'](bold=True, color=WORKBOOK_COLORS['ink'])
        ws.column_dimensions['A'].width = 16
        ws.column_dimensions['B'].width = 42
        _write_table(ws, work_headers[:15], [{k: row[k] for k in work_headers[:15]} for row in works], styles=ops, freeze='A6', start_row=6, sheet_kind='helper')

    unpublished_ws = wb.create_sheet('UNPUBLISHED')
    _write_table(unpublished_ws, work_headers[:15], [{k: row[k] for k in work_headers[:15]} for row in _unpublished_rows()], styles=ops, sheet_kind='helper')

    missing_ws = wb.create_sheet('MISSING_METADATA')
    _write_table(missing_ws, ['work_id', 'series_slug', 'missing_or_weak'], _missing_metadata_rows(), styles=ops, sheet_kind='report')

    validation_ws = wb.create_sheet('VALIDATION_REPORT')
    _write_table(validation_ws, ['level', 'area', 'id', 'issue'], _validation_rows(), styles=ops, sheet_kind='report')

    system_ws = wb.create_sheet('_SYSTEM')
    system_ws['A1'] = 'version'
    system_ws['B1'] = WORKBOOK_VERSION
    system_ws['A2'] = 'exported_at'
    system_ws['B2'] = datetime.now(timezone.utc).isoformat()
    system_ws['A3'] = 'project_root'
    system_ws['B3'] = str(ROOT)
    system_ws.sheet_state = 'hidden'

    EXPORT_DIR.mkdir(parents=True, exist_ok=True)
    final_path = Path(output_path) if output_path else EXPORT_DIR / f'stillmrk-site-workbook-{_now_stamp()}.xlsx'
    wb.save(final_path)
    return final_path


def _sheet_to_dict_rows(ws) -> list[dict[str, Any]]:
    rows = list(ws.iter_rows(values_only=True))
    if not rows:
        return []
    headers = [str(h) if h is not None else '' for h in rows[0]]
    out: list[dict[str, Any]] = []
    for values in rows[1:]:
        if values is None:
            continue
        if all(v is None or v == '' for v in values):
            continue
        out.append({headers[i]: values[i] for i in range(min(len(headers), len(values)))})
    return out


def _parse_jsonish(value: Any) -> Any:
    if value in {None, ''}:
        return []
    if isinstance(value, (list, dict, bool, int, float)):
        return value
    text = str(value)
    try:
        return json.loads(text)
    except json.JSONDecodeError:
        return [item.strip() for item in text.split('|') if item.strip()]


def _norm_compare(value: Any) -> str:
    if value is None:
        return ''
    if isinstance(value, bool):
        return 'true' if value else 'false'
    text = str(value).strip()
    if text.lower() in {'true', 'false'}:
        return text.lower()
    return text


def analyze_workbook_import(workbook_path: str | Path) -> dict[str, Any]:
    ops = _ensure_openpyxl()
    load_workbook = ops['load_workbook']
    wb = load_workbook(filename=workbook_path)
    if '_SYSTEM' in wb.sheetnames and wb['_SYSTEM']['B1'].value != WORKBOOK_VERSION:
        raise RuntimeError('Workbook version is not supported by this importer.')
    summary: dict[str, Any] = {'pages_changed': [], 'works_changed': 0, 'series_changed': 0, 'site_setting_changes': 0, 'documents_changed': 0, 'navigation_changed': 0, 'home_order_changed': 0, 'series_moves': 0, 'unknown_series_slugs': [], 'warnings': []}

    # site settings diff count
    site_rows = _sheet_to_dict_rows(wb['SITE_SETTINGS']) if 'SITE_SETTINGS' in wb.sheetnames else []
    current_site_rows = {(r['group'], r['key']): str(r['value']) for r in _site_settings_rows()}
    for row in site_rows:
        key = (str(row.get('group') or ''), str(row.get('key') or ''))
        if key in current_site_rows and _norm_compare(row.get('value')) != _norm_compare(current_site_rows[key]):
            summary['site_setting_changes'] += 1

    # page diffs
    for page_key in available_page_keys():
        sheet_name = f'PAGE__{page_key}'
        if sheet_name not in wb.sheetnames:
            continue
        incoming_rows = _sheet_to_dict_rows(wb[sheet_name])
        incoming = _unflatten_rows(incoming_rows)
        current = load_page_payload(page_key)
        if page_key == 'home':
            incoming_fs = (incoming.get('featured_series') or {}) if isinstance(incoming, dict) else {}
            current_fs = (current.get('featured_series') or {}) if isinstance(current, dict) else {}
            incoming_fs['series_slugs'] = current_fs.get('series_slugs', [])
            incoming_sw = (incoming.get('selected_works') or {}) if isinstance(incoming, dict) else {}
            current_sw = (current.get('selected_works') or {}) if isinstance(current, dict) else {}
            incoming_sw['work_ids'] = current_sw.get('work_ids', [])
        if incoming != current:
            summary['pages_changed'].append(page_key)

    series_rows = _sheet_to_dict_rows(wb['SERIES_MASTER']) if 'SERIES_MASTER' in wb.sheetnames else []
    current_series = {row['series_slug']: row for row in _series_rows()}
    for row in series_rows:
        slug = str(row.get('series_slug') or '')
        if not slug:
            continue
        if slug not in current_series or any(_norm_compare(row.get(k, '')) != _norm_compare(current_series[slug].get(k, '')) for k in current_series.get(slug, {}).keys()):
            summary['series_changed'] += 1

    works_rows = _sheet_to_dict_rows(wb['WORKS_MASTER']) if 'WORKS_MASTER' in wb.sheetnames else []
    current_works = {row['work_id']: row for row in _work_rows()}
    known_series = set(available_series_slugs())
    unknown_series_slugs: set[str] = set()
    for row in works_rows:
        work_id = str(row.get('work_id') or '')
        if not work_id:
            continue
        incoming_series = str(row.get('series_slug') or '').strip()
        if incoming_series and incoming_series not in known_series:
            unknown_series_slugs.add(incoming_series)
        if work_id not in current_works:
            summary['warnings'].append(f'Unknown work id in workbook: {work_id}')
            continue
        current_series = str(current_works[work_id].get('series_slug') or '').strip()
        if incoming_series and incoming_series != current_series:
            summary['series_moves'] += 1
        if any(_norm_compare(row.get(k, '')) != _norm_compare(current_works[work_id].get(k, '')) for k in current_works[work_id].keys() if k in row):
            summary['works_changed'] += 1
    if unknown_series_slugs:
        summary['unknown_series_slugs'] = sorted(unknown_series_slugs)
        summary['warnings'].append('Unknown series slug(s) in workbook: ' + ', '.join(sorted(unknown_series_slugs)))

    home_rows = _sheet_to_dict_rows(wb['HOME_ORDER']) if 'HOME_ORDER' in wb.sheetnames else []
    current_home_rows = _home_order_rows()
    if [(r.get('group'), r.get('item_id'), r.get('position'), r.get('enabled')) for r in home_rows] != [(r.get('group'), r.get('item_id'), r.get('position'), r.get('enabled')) for r in current_home_rows]:
        summary['home_order_changed'] = 1

    doc_rows = _sheet_to_dict_rows(wb['DOCUMENTS']) if 'DOCUMENTS' in wb.sheetnames else []
    if len(doc_rows) != len(_document_rows()) or any(_norm_compare(a.get(k,'')) != _norm_compare(b.get(k,'')) for a, b in zip(doc_rows, _document_rows()) for k in ['id','title','kind','description','file','audience','featured']):
        summary['documents_changed'] = 1

    nav_rows = _sheet_to_dict_rows(wb['NAVIGATION']) if 'NAVIGATION' in wb.sheetnames else []
    if len(nav_rows) != len(_navigation_rows()) or any(_norm_compare(a.get(k,'')) != _norm_compare(b.get(k,'')) for a, b in zip(nav_rows, _navigation_rows()) for k in ['position','label','href','page','visible']):
        summary['navigation_changed'] = 1
    return summary


def import_site_workbook(workbook_path: str | Path) -> dict[str, Any]:
    ops = _ensure_openpyxl()
    load_workbook = ops['load_workbook']
    wb = load_workbook(filename=workbook_path)
    if '_SYSTEM' in wb.sheetnames and wb['_SYSTEM']['B1'].value != WORKBOOK_VERSION:
        raise RuntimeError('Workbook version is not supported by this importer.')
    applied = {'pages': 0, 'series': 0, 'works': 0, 'site_settings': 0, 'documents': 0, 'navigation': 0, 'home': 0, 'derivatives_regenerated': 0}

    # Site settings
    site_yaml = load_yaml(CONTENT_DIR / 'site.yaml') or {}
    artist_yaml = load_yaml(CONTENT_DIR / 'artist.yaml') or {}
    for row in _sheet_to_dict_rows(wb['SITE_SETTINGS']):
        group = str(row.get('group') or '')
        key = str(row.get('key') or '')
        value = row.get('value')
        value_type = str(row.get('type') or 'string')
        target = site_yaml if group == 'site' else artist_yaml if group == 'artist' else None
        if target is None or not key:
            continue
        parsed = _parse_jsonish(value) if value_type == 'json' else ((str(value).strip().lower() in {'true','1','yes','y'}) if value_type == 'bool' else value)
        target[key] = parsed
    write_yaml(CONTENT_DIR / 'site.yaml', site_yaml)
    write_yaml(CONTENT_DIR / 'artist.yaml', artist_yaml)
    applied['site_settings'] = 1

    # Page sheets
    for page_key in available_page_keys():
        sheet_name = f'PAGE__{page_key}'
        if sheet_name not in wb.sheetnames:
            continue
        incoming_rows = _sheet_to_dict_rows(wb[sheet_name])
        incoming_payload = _unflatten_rows(incoming_rows)
        current_payload = load_page_payload(page_key)
        if page_key == 'home':
            current_fs = current_payload.get('featured_series') if isinstance(current_payload.get('featured_series'), dict) else {}
            current_sw = current_payload.get('selected_works') if isinstance(current_payload.get('selected_works'), dict) else {}
            inc_fs = incoming_payload.get('featured_series') if isinstance(incoming_payload.get('featured_series'), dict) else {}
            inc_sw = incoming_payload.get('selected_works') if isinstance(incoming_payload.get('selected_works'), dict) else {}
            inc_fs['series_slugs'] = current_fs.get('series_slugs', [])
            inc_sw['work_ids'] = current_sw.get('work_ids', [])
            incoming_payload['featured_series'] = inc_fs
            incoming_payload['selected_works'] = inc_sw
        save_page_payload(page_key, incoming_payload)
        applied['pages'] += 1

    # Documents
    doc_rows = _sheet_to_dict_rows(wb['DOCUMENTS']) if 'DOCUMENTS' in wb.sheetnames else []
    resources = {'downloads': []}
    for row in doc_rows:
        resources['downloads'].append({
            'id': str(row.get('id') or ''),
            'title': row.get('title') or '',
            'description': row.get('description') or '',
            'file': row.get('file') or '',
            'kind': row.get('kind') or '',
            'audience': row.get('audience') or '',
            'featured': bool(row.get('featured', False)),
        })
    write_yaml(CONTENT_DIR / 'resources.yaml', resources)
    applied['documents'] = len(doc_rows)

    # Navigation
    nav_rows = sorted(_sheet_to_dict_rows(wb['NAVIGATION']) if 'NAVIGATION' in wb.sheetnames else [], key=lambda row: int(row.get('position') or 9999))
    navigation = {'items': []}
    for row in nav_rows:
        navigation['items'].append({
            'label': row.get('label') or '',
            'href': row.get('href') or '',
            'page': row.get('page') or '',
            'visible': bool(row.get('visible', True)),
        })
    write_yaml(CONTENT_DIR / 'navigation.yaml', navigation)
    applied['navigation'] = len(nav_rows)

    # Series master update
    series_rows = sorted(_sheet_to_dict_rows(wb['SERIES_MASTER']) if 'SERIES_MASTER' in wb.sheetnames else [], key=lambda row: int(row.get('order') or 9999))
    for row in series_rows:
        slug = str(row.get('series_slug') or '').strip()
        if not slug:
            continue
        path = series_file_for_slug(slug)
        if path.exists():
            payload = load_yaml(path) or {}
        else:
            create_series_file(slug, str(row.get('title') or slug), str(row.get('years') or ''), str(row.get('mood') or ''), str(row.get('description') or ''))
            path = series_file_for_slug(slug)
            payload = load_yaml(path) or {}
        payload['slug'] = slug
        payload['title'] = row.get('title') or ''
        payload['years'] = row.get('years') or ''
        payload['mood'] = row.get('mood') or ''
        payload['description'] = row.get('description') or ''
        payload['cover_work_id'] = row.get('cover_work_id') or ''
        payload['order'] = int(row.get('order') or 0)
        payload['visibility'] = row.get('visibility') or 'public'
        payload['review_mode'] = bool(row.get('review_mode', False))
        payload['allow_favorites'] = bool(row.get('allow_favorites', True))
        payload['allow_inquiry_basket'] = bool(row.get('allow_inquiry_basket', True))
        payload['project_type'] = row.get('project_type') or ''
        payload['download_ids'] = [item.strip() for item in str(row.get('download_ids') or '').split('|') if item.strip()]
        write_yaml(path, payload)
        applied['series'] += 1

    # Works and series positions
    pipeline = load_pipeline()
    current_map = work_to_series_map()
    grouped: dict[str, list[tuple[int, str]]] = defaultdict(list)
    work_rows = _sheet_to_dict_rows(wb['WORKS_MASTER']) if 'WORKS_MASTER' in wb.sheetnames else []
    for row in work_rows:
        work_id = str(row.get('work_id') or '').strip()
        if not work_id:
            continue
        path = work_file_for_id(work_id)
        if not path.exists():
            continue
        payload = load_yaml(path) or {}
        old_series = str(payload.get('series') or current_map.get(work_id) or '').strip()
        new_series = str(row.get('series_slug') or old_series).strip()
        payload['title'] = row.get('title') or ''
        payload['series'] = new_series
        payload['published'] = bool(row.get('published', False))
        payload['year'] = row.get('year') or ''
        payload['location'] = row.get('location') or ''
        payload['alt'] = row.get('alt') or ''
        payload['caption'] = row.get('caption') or ''
        payload['tags'] = [item.strip() for item in str(row.get('tags') or '').split('|') if item.strip()]
        payload['hero_safe'] = bool(row.get('hero_safe', False))
        payload['grid_safe'] = bool(row.get('grid_safe', False))
        payload['social_safe'] = bool(row.get('social_safe', False))
        payload['focal_point'] = {'x': int(row.get('focal_x') or 50), 'y': int(row.get('focal_y') or 50)}
        payload['project_type'] = row.get('project_type') or ''
        payload['licensing_available'] = bool(row.get('licensing_available', False))
        payload['print_available'] = bool(row.get('print_available', False))
        payload['price_note'] = row.get('price_note') or ''
        image = payload.get('image') if isinstance(payload.get('image'), dict) else {}
        image['master'] = row.get('image_master') or image.get('master') or ''
        image['render_name'] = row.get('render_name') or work_id
        payload['image'] = image
        if old_series and new_series and old_series != new_series:
            moved_source = move_original_between_series(work_id, old_series, new_series, pipeline=pipeline)
            move_generated_between_series(work_id, old_series, new_series, pipeline=pipeline)
            source_for_derivatives = moved_source or source_path_for_work(new_series, work_id, pipeline=pipeline)
            if source_for_derivatives and source_for_derivatives.exists():
                generate_derivatives(source_for_derivatives, new_series, image.get('render_name') or work_id, pipeline=pipeline, force=True)
                applied['derivatives_regenerated'] += 1
        write_yaml(path, payload)
        grouped[new_series].append((int(row.get('series_position') or 9999), work_id))
        applied['works'] += 1

    # Apply grouped order back to series files.
    for slug in available_series_slugs():
        path = series_file_for_slug(slug)
        payload = load_yaml(path) or {}
        ordered = [work_id for _, work_id in sorted(grouped.get(slug, []), key=lambda item: (item[0], item[1]))]
        if ordered:
            payload['work_ids'] = ordered
            if str(payload.get('cover_work_id') or '').strip() not in ordered:
                payload['cover_work_id'] = ordered[0]
            write_yaml(path, payload)

    # Home order
    home_rows = sorted(_sheet_to_dict_rows(wb['HOME_ORDER']) if 'HOME_ORDER' in wb.sheetnames else [], key=lambda row: (str(row.get('group') or ''), int(row.get('position') or 9999)))
    featured_series = [str(row.get('item_id') or '') for row in home_rows if str(row.get('group') or '') == 'featured_series' and bool(row.get('enabled', True)) and str(row.get('item_id') or '').strip()]
    selected_works = [str(row.get('item_id') or '') for row in home_rows if str(row.get('group') or '') == 'selected_works' and bool(row.get('enabled', True)) and str(row.get('item_id') or '').strip()]
    reorder_homepage_featured(featured_series=featured_series, selected_works=selected_works)
    applied['home'] = len(home_rows)

    ensure_og_images_from_content(load_content_for_og(), force=True)
    return applied
